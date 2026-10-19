from pathlib import Path
from bs4 import BeautifulSoup
from openpyxl import load_workbook
import gzip,base64,json,re,html
R=Path('/home/user');P=R/'UK_Junior_Data_Analyst_Roadmap_Eastbourne.html'
backup=R/'reference/roadmap_v5_before_copy_cleanup.html.gz'
if not backup.exists():backup.write_bytes(gzip.compress(P.read_bytes()))
s=BeautifulSoup(gzip.decompress(backup.read_bytes()).decode(),'html.parser')
def sethtml(node,text):
 node.clear()
 for x in list(BeautifulSoup(text,'html.parser').contents):node.append(x)
def title(id,h,p,k=None):
 t=s.find(id=id).select_one('.section-title');sethtml(t.h2,h);sethtml(t.p,p)
 if k:sethtml(t.select_one('.kicker'),k)
# Reader-facing copy only: the curriculum, art, layout and stable progress keys are unchanged.
s.title.string='Operation Data Analyst · UK South East Roadmap'
s.select_one('.hero-eyebrow').contents[-1].replace_with(' A PRACTICAL DATA ANALYTICS ROADMAP / 2026')
sethtml(s.select_one('.hero .lead'),'Build job-ready analytics skills alongside work, study or other commitments. A flexible roadmap for <b>South-East London, Brighton and the wider South East.</b> Excel, SQL and Power BI first. Turn practice into a portfolio you can explain.')
s.select_one('.estimate').string='9 training arcs · 142 checkpoints · 5 projects · free core learning resources.'
title('training','Every session makes you stronger.','Start small. Practise the fundamentals. Keep going when a task takes more than one attempt.')
for p in s.select('#training .arc-card p'):
 if 'Rock Lee' in p.get_text():p.string='Rock Lee trained the fundamentals. Start where you are and build through repetition—one checked query, one clearer explanation, one stronger skill at a time.'
title('motivation','Keep showing up.','You don’t need to master everything today. Finish one useful task, understand it, and return for the next.')
# Keep useful evidence; remove the narration of the website-building and agent-review process.
title('audit','What employers look for.','Explore example roles, the skills they ask for and the experience needed. Check each advert’s status before applying.','ROLE EVIDENCE / 18 OCT 2026')
s.select_one('#audit .audit-summary .audit-tag').string='SKILLS THAT MATTER'
s.select_one('#audit .audit-summary h3').string='Focus on useful, demonstrable skills.'
s.select_one('#audit .audit-summary p').string='Clean data, reliable queries, clear reports and sound judgement. Build evidence that shows how you work—not just a list of courses you have watched.'
for el in list(s.select('#audit .review-lenses')):el.decompose()
for p in list(s.select('#audit p')):
 if 'same assistant' in p.get_text() or 'independent agent partners' in p.get_text():p.decompose()
for p in s.select('#market .intel-card p'):
 if 'now explicit checkpoints' in p.get_text():p.string='A useful report needs more than charts. Learn refresh, workspace access and row-level security alongside data preparation, modelling and DAX.'
for t in s.select('#market .audit-tag'):
 text=t.get_text();t.string={'Keep as core':'Core skill','Keep, then translate':'Core + dialect bridge','Keep and extend':'Core + report operations'}.get(text,text)
# Current guidance rather than a history of earlier recommendations.
title('courses','Learn with purpose.','Choose one main course per tool. Use focused lessons to fill gaps, then practise without the tutorial.','LEARNING RESOURCES')
for dt in s.select('.course-card dt'):
 if dt.get_text()=='Verified or inspected scope':dt.string='What to expect'
for c in s.select('.course-card'):
 if c.select_one('h3') and 'SQL foundations' in c.select_one('h3').get_text():
  c.select_one('.audit-tag').string='FOUNDATIONS + FOCUSED PRACTICE'
  c.select_one('dd').string='Start with analyst-focused SQL foundations: filtering, joins, aggregations, CASE and CTEs. Use selected intermediate lessons for window functions, dates and data cleaning, then practise translating queries to T-SQL.'
# Retain the syllabus coverage map, but describe a curriculum rather than a changelog.
title('coverage','The skills to practise—and how to prove them.','Use this checklist to connect each skill with a practical piece of work. Keep role-specific additions separate from the core path.','SYLLABUS AT A GLANCE')
coverage=s.select_one('#coverage table')
coverage.select('th')[1].string='Study focus'
for row in coverage.select('tbody tr'):
 cell=row.find_all('td')[1];v=cell.get_text()
 cell.string={'Keep':'Core practice','Make explicit':'Practical application','Fix resource mismatch':'Foundations + targeted lessons','Add short bridge':'Short dialect bridge','Add explicit tasks':'Hands-on practice','Keep and apply':'Apply to real questions','Keep and strengthen':'Reproducible practice','Targeted optional pack':'Role-specific study','Targeted practice':'Domain practice','Optional awareness':'Optional awareness','Deliberately excluded':'Outside the core path'}.get(v,v)
note=s.select_one('#coverage .note')
sethtml(note,'<b>Use the syllabus alongside job criteria:</b> this roadmap covers core junior reporting and analytics skills. Employers may also require domain knowledge, specific systems or prior experience. Check the essential criteria for each role. Certification preparation has its own syllabus.')
title('roadmap','Your 28-week roadmap.','Practise each skill, then complete the boss task. Select the course chapters you need and spend the rest of the session building, checking and explaining your work.')
s.select_one('.roadmap-index p').string='Tick a skill when you can demonstrate it without copying. Move forward when the boss task is ready.'
# Flexible enough for another learner to use without inheriting someone else’s employment history.
replacements={
'Start where you are: Excel experience and early SQL practice. This is a diagnostic and setup week, not a restart from zero.':'Check your starting point, set up your tools and choose a regular study routine. Use the diagnostic to identify what you already know and what needs practice.',
'Use only public or synthetic data; never upload Lidl/customer/colleague data.':'Use only public or synthetic practice data; never upload confidential employer, customer or colleague information.',
'Set a repeatable 12-hour weekly schedule around your actual rota.':'Set a repeatable 12-hour weekly schedule around your work, study and other commitments.',
'Prepare a truthful 60-second Lidl → analytics career-change story.':'Prepare a truthful 60-second introduction connecting your background, analytics practice and target role.',
'Only describe yourself as ex-Lidl if you have actually left; otherwise list your current role.':'Describe your employment status accurately; list a role as current only while you still hold it.',
'A schedule that survives Lidl shifts.':'A routine that fits real life.',
'These replace the PDF’s web-development apps with pure analytics deliverables. No API, React or full-stack requirement.':'Build from a clean spreadsheet to an end-to-end analysis. Each project answers a clear question and gives you something useful to discuss in an interview.',
'Reuse public/synthetic retail data; do not present it as Lidl data':'Reuse public or synthetic retail data and label its source clearly',
'project outcomes are not workplace achievements. Say “built a synthetic retail report with X validated measures”, not “saved Lidl £X”. Only claim time savings you actually measured in a real process.':'project outcomes are not workplace achievements. Describe the report you built, the checks you performed and the findings you can verify. Only claim savings or workplace impact you actually measured.',
'Transfer retail skills honestly:':'Show transferable skills honestly:',
'Current employment wording:':'Keep employment details accurate:',
'do not say “ex-Lidl” unless you have left.':'use truthful role titles, employment dates and responsibilities.',
'1. Why move from Lidl into analytics?':'1. Why do you want to work in data analytics?',
'40 questions with answer directions. Speak first, then reveal the prompt. Use your own examples rather than memorising a script.':'40 questions with answer guidance. Answer aloud before opening the guidance, and use your own examples rather than memorising a script.',
'An explicit self-review of the plan. This is not a claim that an independent second agent reviewed it.':'Protect your time, avoid tutorial overload and build work you can trust.',
'07 · Retail experience can be oversold.':'07 · Experience can be oversold.',
'Use truthful transfer skills and clearly label public/synthetic projects. No invented internal Lidl data, salaries, savings or interview rates.':'Describe your experience honestly and label public or synthetic projects clearly. Use numbers only when you can explain how they were measured.',
'The revised plan explicitly covers the junior reporting/analytics skill groups identified in this sample.':'This roadmap focuses on core junior reporting and analytics skills.',
'Useful reference; syntax can differ from the MySQL course.':'SQL reference; check syntax against the database you are using.',
'No JavaScript, web frameworks, Business Analyst core or machine-learning syllabus.':'Leave unrelated tools and advanced specialisms until your target roles justify them.'
}
for node in list(s.find_all(string=True)):
 if node.parent.name in ['script','style']:continue
 text=str(node)
 for old,new in replacements.items():text=text.replace(old,new)
 if text!=str(node):node.replace_with(text)
for code in s.select('#jobhunt .codebox'):
 if 'Hi [Name]' in code.get_text():code.string='Hi [Name], I’m building towards a junior data-analytics role, with a background in [your current work or study].\n\nI’ve built [project + link] using [Excel / SQL / Power BI]. It answers [specific question], with [one verified finding or quality check].\n\nI saw [role + reference] at [company]. If you feel my work is a fit, would you be open to a short conversation or pointing me towards the right application route?\n\nThanks for your time,\n[Your name]'
for el in list(s.select('#review .note')):
 if 'Audit verdict' in el.get_text():el.decompose()
# Sources should help a learner find material, not narrate the authoring process.
source=s.find(id='sources');lst=source.select_one('.source-list');lst.clear()
resources=[('https://www.prospects.ac.uk/job-profiles/data-analyst/','Prospects · Data analyst role, skills and salary','UK career guidance; salary figures are national indicators, not local guarantees.'),('https://www.youtube.com/watch?v=pCJ15nGFgVg','Excel for Data Analytics · Luke Barousse','Main Excel course; select chapters for the skill you are practising.'),('https://www.youtube.com/watch?v=7mz73uXD9DA','SQL for Data Analytics · Luke Barousse','Analyst-focused SQL foundations and project.'),('https://www.youtube.com/watch?v=QKIGsShyEsQ','Intermediate SQL · Luke Barousse','Selected window-function, date and cleaning lessons.'),('https://www.youtube.com/watch?v=FwjaHCVNBWA','Power BI for Data Analytics · Luke Barousse','Data preparation, modelling, DAX and dashboard practice.'),('https://www.youtube.com/watch?v=wUSDVGivd-8','Python for Data Analytics · Luke Barousse','Python and pandas practice; use current course notebooks.')]
for u,t,d in resources:
 li=s.new_tag('li');sethtml(li,f'<a href="{u}" target="_blank" rel="noopener noreferrer">{t} ↗</a> — {d}');lst.append(li)
panel=source.select_one('.panel');sethtml(panel,'<h3>Using this roadmap</h3><p>Adapt the pace to your experience and commitments. This is a learning plan, not a guarantee of employment. Check current role requirements, salaries and application deadlines before applying.</p><p style="margin-top:12px">Save the HTML to use the roadmap offline. External videos and documentation require internet access. Progress is stored in your browser; export a backup before switching devices. The PDF is a static reading copy, and Excel progress is tracked separately.</p><p class="muted" style="margin-top:12px">Anime artwork belongs to its respective rights holders. Reuse is subject to those rights.</p>')
details=source.select_one('.audit-details');summary=details.find('summary');sethtml(summary,'Additional references <span>Documentation &amp; course files</span>')
body=details.select_one('.evidence-body');sethtml(body,'<p>Use official documentation for current features, access controls and exam requirements. Course interfaces and files can change; use the latest materials from the course provider.</p>')
ev=json.loads((R/'reference/audit_evidence.json').read_text());ul=s.new_tag('ul')
notes={
'Excel course outline':'Excel lesson files and chapter list. Use the repository ZIP or alternate file link if an individual workbook fails to open.',
'SQL foundations':'PostgreSQL setup, free starter files and project resources. Course datasets are historical, not current hiring evidence.',
'Intermediate SQL':'Window-function and analysis practice, with a fictional retail dataset.',
'Power BI course outline':'Power Query, DAX and project files. Use the provider’s complete download link for large files.',
'Python course outline':'Current notebooks for pandas, cleaning and visualisation.',
'Current PL-300 study guide':'The official exam objectives. Review them separately if you choose certification.',
'T-SQL learning path':'SQL Server querying fundamentals for the short dialect bridge.',
'T-SQL window lesson':'OVER, ranking, aggregate and offset functions.',
'UK GDPR principles':'Official guidance on purpose, minimisation, accuracy, security and accountability.'}
for d in ev['course_sources']:
 if d['name'] not in notes:continue
 li=s.new_tag('li');sethtml(li,f'<a href="{html.escape(d["url"])}" target="_blank" rel="noopener noreferrer">{d["name"]} ↗</a> — {notes[d["name"]]}');ul.append(li)
body.append(ul)
# Footer and accessibility text should also stand alone.
footer=s.find('footer');footer.find('h3').string='Operation Data Analyst · UK South East'
footer.find_all('p')[0].string='Ask a useful question. Check the data. Explain the answer.'
footer.find_all('p')[-1].string='October 2026 · Save the HTML for offline use · Keep a progress backup.'
for im in s.find_all('img'):
 alt=im.get('alt','')
 if 'original artwork extracted' in alt:im['alt']=alt.split(' — ')[0]+' motivational artwork'
 elif 'reference PDF' in alt:im['alt']='Learner looking towards a city on the horizon'
# Navigation and course labels should read like learning guidance, not editorial notes.
for a in s.select('nav a[href="#audit"]'):a.string='Roles'
for a in s.select('.hero-badges a'):a.string='UK skills & role guidance ↗'
labels={'KEEP · SELECT CHAPTERS':'EXCEL · CORE','KEEP + OPERATIONS PACK':'POWER BI · CORE','KEEP · LATER IN SEQUENCE':'PYTHON · AFTER THE CORE','KEEP + APPLIED TASKS':'STATISTICS · APPLIED'}
for tag in s.select('.course-card .audit-tag'):
 text=tag.get_text();tag.string=labels.get(text,text)
# A course decision table is still useful; construction-history wording is not.
for node in list(s.find_all(string=True)):
 if node.parent.name in ['style','script']:continue
 t=str(node).replace('Review decision','Study focus').replace('What changes for you:','What this means for your study:').replace(' No BA core.','').replace('direct fetch reached verification screen','employer page access limited').replace('Direct fetch reached verification screen','Employer page access limited').replace('direct fetch confirmed vacancy-not-found redirect','employer page no longer available')
 if t!=str(node):node.replace_with(t)
# Synchronise learner-facing checkpoint descriptions without changing keys or counts.
keys={b['data-key']:b.find_parent('label').find('span').get_text() for b in s.select('.checklist input[data-key]')}
wb=load_workbook(R/'Eastbourne_Data_Analyst_Tracker.xlsx')
for row in wb['Checkpoints'].iter_rows(min_row=2):
 if row[0].value in keys:row[3].value=keys[row[0].value]
wb['Start here']['A1']='OPERATION DATA ANALYST · UK SOUTH EAST'
for row in wb['Coverage audit'].iter_rows(min_row=2):
 v=row[1].value
 row[1].value={'Keep':'Core practice','Make explicit':'Practical application','Fix resource mismatch':'Foundations + targeted lessons','Add short bridge':'Short dialect bridge','Add explicit tasks':'Hands-on practice','Keep and apply':'Applied practice','Keep and strengthen':'Reproducible practice'}.get(v,v)
wb.save(R/'Eastbourne_Data_Analyst_Tracker.xlsx')
s.select_one('a[download="Eastbourne_Data_Analyst_Tracker.xlsx"]')['href']='data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,'+base64.b64encode((R/'Eastbourne_Data_Analyst_Tracker.xlsx').read_bytes()).decode()
for filename in ['arc_data_v5.json']:
 data=json.loads((R/'reference'/filename).read_text())
 for a in data:
  a['items']=[keys[f'{a["id"]}-{i}'] for i in range(len(a['items']))]
 (R/'reference'/filename).write_text(json.dumps(data,indent=2))
P.write_text(str(s))
visible=BeautifulSoup(str(s),'html.parser')
for x in visible(['script','style']):x.decompose()
text=visible.get_text('\n',strip=True);(R/'reference/visible_copy_cleaned.txt').write_text(text)
for phrase in ['your PDF','reference PDF','same assistant','agent partners','independent second agent','earlier core stack','saved ticks retain their keys','Kerala reference','Lidl']:
 assert phrase.casefold() not in text.casefold(),phrase
assert len(keys)==142
print('Copy cleanup complete. Preserved',len(keys),'checkpoint keys and',len(s.find_all('img')),'images.')
